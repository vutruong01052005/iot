from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


OUTPUT = Path(__file__).with_name("Bao_cao_Bai_2_IoT.docx")
NAVY = "17324D"
TEAL = "087E8B"
PALE = "EAF2F5"
WHITE = "FFFFFF"


def set_cell_fill(cell, color):
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    properties.append(shading)


def set_cell_text(cell, text, bold=False, color=None, size=9):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(str(text))
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(document, headers, rows, widths=None):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    for index, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], header, True, WHITE, 9)
        set_cell_fill(table.rows[0].cells[index], NAVY)
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(row):
            set_cell_text(cells[index], value)
            if row_index % 2 == 1:
                set_cell_fill(cells[index], PALE)
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Cm(width)
    document.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_bullets(document, items):
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.add_run(item)


def add_numbered(document, items):
    for item in items:
        paragraph = document.add_paragraph(style="List Number")
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.add_run(item)


def add_code(document, text):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(0.35)
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(7)
    run = paragraph.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string("24445C")


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Trang ")
    run.font.name = "Arial"
    run.font.size = Pt(9)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def setup_document(document):
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.1)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.1)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")

    for style_name, size, color in [
        ("Title", 24, NAVY),
        ("Heading 1", 16, NAVY),
        ("Heading 2", 12, TEAL),
    ]:
        style = document.styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_run = header.add_run("BÀI THỰC HÀNH 2  |  IoT VÀ ỨNG DỤNG")
    header_run.font.name = "Arial"
    header_run.font.size = Pt(8)
    header_run.font.color.rgb = RGBColor.from_string(TEAL)
    add_page_number(section.footer.paragraphs[0])


def add_paragraph(document, text):
    paragraph = document.add_paragraph(text)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return paragraph


def build_report():
    document = Document()
    setup_document(document)

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(68)
    run = paragraph.add_run("BÀI THỰC HÀNH MÔN HỌC\nIoT VÀ ỨNG DỤNG")
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(17)
    run.font.color.rgb = RGBColor.from_string(NAVY)

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(38)
    run = paragraph.add_run("BÁO CÁO BÀI THỰC HÀNH SỐ 2")
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(21)
    run.font.color.rgb = RGBColor.from_string(TEAL)

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(10)
    run = paragraph.add_run("THU THẬP, LƯU TRỮ VÀ\nTIỀN XỬ LÝ DỮ LIỆU IoT")
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor.from_string(NAVY)

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(18)
    paragraph.add_run("ESP32/Wokwi → MQTT → InfluxDB → Streamlit").italic = True

    document.add_paragraph().paragraph_format.space_after = Pt(45)
    add_table(
        document,
        ["Thông tin sinh viên", "Nội dung"],
        [
            ("Họ và tên", "Vũ Xuân Trường"),
            ("Mã sinh viên", "B23DCAT312"),
            ("Lớp / nhóm thực hành", "N15"),
            ("Ngày thực hiện", "15/09/2026"),
        ],
        [6.0, 9.0],
    )
    paragraph = document.add_paragraph("Năm học 2026")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_page_break()

    document.add_heading("Tóm tắt", level=1)
    add_paragraph(
        document,
        "Báo cáo trình bày pipeline IoT kế thừa node cảm biến ESP32 của Bài thực hành 1. "
        "ESP32 trên Wokwi đọc nhiệt độ, độ ẩm, ánh sáng và khoảng cách; dữ liệu JSON được "
        "publish qua MQTT, collector Python ghi vào InfluxDB, sau đó script tiền xử lý tạo "
        "các đặc trưng và dashboard Streamlit phục vụ giám sát.",
    )
    add_paragraph(
        document,
        "Trong lần kiểm tra hệ thống, PlatformIO build thành công; serial ghi nhận publish=OK; "
        "InfluxDB lưu được các field cảm biến; script tiền xử lý xử lý 118 cửa sổ dữ liệu và "
        "dashboard trả HTTP 200. Đây là kết quả kiểm tra chức năng trên môi trường lab, "
        "không phải phép đo benchmark hiệu năng.",
    )

    document.add_heading("1. Giới thiệu", level=1)
    add_paragraph(
        document,
        "IoT kết nối cảm biến, thiết bị xử lý và nền tảng phần mềm để thu thập, truyền, lưu trữ "
        "và phân tích dữ liệu. Bài thực hành 1 tập trung vào node ESP32 và truyền telemetry. "
        "Bài thực hành 2 mở rộng quy trình đó bằng broker MQTT, collector, cơ sở dữ liệu "
        "time-series, bước tiền xử lý và giao diện giám sát.",
    )
    add_paragraph(
        document,
        "Pipeline được triển khai cục bộ bằng Docker Compose; firmware được build bằng PlatformIO "
        "và mô phỏng trong Wokwi. Dữ liệu raw được giữ riêng với dữ liệu sau xử lý để có thể "
        "đối chiếu và chạy lại quy trình phân tích.",
    )

    document.add_heading("2. Mục tiêu", level=1)
    add_bullets(
        document,
        [
            "Thu thập telemetry cảm biến qua MQTT và parse JSON ở collector Python.",
            "Lưu dữ liệu theo thời gian trong InfluxDB 2.x.",
            "Xử lý missing values và outlier, resample chuỗi thời gian, tạo rolling mean và delta.",
            "Chuẩn hóa dữ liệu bằng Min-Max và hiển thị raw/processed trên dashboard.",
            "Kiểm tra đường đi dữ liệu từ ESP32 đến InfluxDB và giao diện Streamlit.",
        ],
    )

    document.add_heading("3. Kiến trúc hệ thống", level=1)
    add_table(
        document,
        ["Tầng", "Thành phần", "Chức năng"],
        [
            ("Thiết bị", "ESP32 / Wokwi", "Đọc cảm biến và publish JSON mỗi 5 giây."),
            ("Truyền thông", "Mosquitto MQTT", "Nhận telemetry trên topic iot/bai2/esp32/telemetry."),
            ("Thu thập", "Python Collector", "Subscribe, parse payload và ghi field số hợp lệ."),
            ("Lưu trữ", "InfluxDB 2.x", "Lưu sensor_raw và sensor_preprocessed trong bucket iot_raw."),
            ("Phân tích", "Preprocess script", "Missing, IQR, resampling, rolling mean, delta, Min-Max."),
            ("Hiển thị", "Streamlit", "Metric mới nhất, biểu đồ raw và dữ liệu đã xử lý."),
        ],
        [2.7, 4.0, 8.3],
    )

    document.add_heading("4. Thiết bị và kết nối cảm biến", level=1)
    add_paragraph(
        document,
        "Firmware được quản lý bằng PlatformIO. Cấu hình chân trong firmware và diagram.json "
        "được đồng bộ như sau:",
    )
    add_table(
        document,
        ["Thiết bị", "Chân ESP32", "Mục đích"],
        [
            ("DHT22", "GPIO15", "Nhiệt độ (°C), độ ẩm (%RH)."),
            ("Photoresistor module", "AO → GPIO34", "Đọc ánh sáng analog và quy đổi sang lux."),
            ("HC-SR04", "TRIG GPIO5, ECHO GPIO18", "Đo khoảng cách (cm)."),
            ("LED", "GPIO2", "Báo trạng thái publish MQTT thành công."),
        ],
        [4.0, 4.2, 6.8],
    )
    add_paragraph(
        document,
        "DHT22 dùng thư viện DHT sensor library for ESPx (DHTesp). LDR được cấp 3.3 V và GND; "
        "ngõ AO nối GPIO34. Thuộc tính lux của linh kiện Wokwi có thể thay đổi trong popup cảm biến. "
        "Firmware dùng ADC 12-bit và lấy trung bình nhiều lần đọc để giảm nhiễu.",
    )

    document.add_heading("5. Truyền dữ liệu MQTT", level=1)
    add_paragraph(
        document,
        "Telemetry được publish lên iot/bai2/esp32/telemetry. Trạng thái thiết bị được publish "
        "lên iot/bai2/esp32/status. Mosquitto chạy trong Docker và publish cổng 1883 lên host.",
    )
    add_code(
        document,
        '{"device_id":"esp32","ts":1790752222040,"seq":1,'
        '"temperature":25.00,"humidity":50.00,"light_lux":499.6,'
        '"distance_cm":100.86,"rssi":-96,"uptime_s":10}',
    )
    add_paragraph(
        document,
        "Các field cảm biến không hợp lệ được serialize thành null. Collector vẫn giữ mẫu và chỉ "
        "ghi các field cảm biến có kiểu số; do đó cần kiểm tra cả payload Serial và measurement "
        "sensor_raw khi DHT trả null.",
    )
    add_paragraph(
        document,
        "Khi ESP32/Wokwi kết nối tới Mosquitto trên máy host, MQTT_SERVER phải là IPv4 của host, "
        "không phải localhost hay hostname Docker mosquitto. IP có thể đổi theo DHCP; Wokwi cần "
        "Private IoT Gateway để truy cập broker cục bộ.",
    )

    document.add_heading("6. Collector và InfluxDB", level=1)
    add_paragraph(
        document,
        "Collector subscribe topic telemetry, kiểm tra các khóa cấu trúc bắt buộc, chuyển các giá "
        "trị số sang float/integer và ghi điểm dữ liệu vào InfluxDB. Field null được bỏ qua; các "
        "field hợp lệ khác của cùng mẫu vẫn được ghi.",
    )
    add_table(
        document,
        ["Bucket", "Measurement", "Nội dung"],
        [
            ("iot_raw", "sensor_raw", "Telemetry raw, tag device và các field cảm biến/diagnostic."),
            ("iot_raw", "sensor_preprocessed", "Dữ liệu resample, rolling mean, delta và Min-Max."),
        ],
        [3.0, 4.0, 8.0],
    )
    add_paragraph(
        document,
        "Docker Compose khởi tạo InfluxDB 2.7, Mosquitto, collector và dashboard. InfluxDB được "
        "cấu hình bằng org iot-lab, bucket iot_raw và token demo trong compose; thông tin này chỉ "
        "dùng cho môi trường học tập cục bộ.",
    )

    document.add_heading("7. Tiền xử lý dữ liệu", level=1)
    add_numbered(
        document,
        [
            "Chuyển các cột cảm biến và RSSI đang có sang kiểu số; xử lý missing bằng nội suy và forward/backward fill.",
            "Tính ngưỡng IQR theo từng cột; giá trị ngoài ngưỡng được đánh dấu missing rồi nội suy lại.",
            "Resample dữ liệu theo cửa sổ 10 giây và lấy trung bình.",
            "Tạo rolling mean 3 mẫu cho nhiệt độ, độ ẩm và ánh sáng; tạo distance_delta cho khoảng cách.",
            "Tạo field Min-Max chuẩn hóa trong khoảng 0–1; nếu min bằng max, giá trị chuẩn hóa được đặt bằng 0.",
            "Ghi kết quả vào measurement sensor_preprocessed. Khi thiếu DHT, script cảnh báo và xử lý các sensor khả dụng thay vì crash.",
        ],
    )

    document.add_heading("8. Dashboard", level=1)
    add_paragraph(
        document,
        "Dashboard Streamlit chạy tại cổng 8501. Trang hiển thị nhiệt độ, độ ẩm, ánh sáng, khoảng "
        "cách; biểu đồ sensor_raw và bảng các mẫu gần nhất; nếu đã chạy preprocess, hiển thị thêm "
        "measurement sensor_preprocessed. Metric ưu tiên mẫu raw mới nhất có giá trị hợp lệ và "
        "fallback sang processed khi raw thiếu field.",
    )

    document.add_heading("9. Quy trình chạy", level=1)
    add_numbered(
        document,
        [
            "Tại repo root, chạy docker compose up -d --build; xác nhận InfluxDB healthy và các dịch vụ khác running bằng docker compose ps.",
            "Mở thư mục firmware trong VS Code; chạy PlatformIO: Build, sau đó Wokwi: Start Simulator.",
            "Giữ tab Wokwi hiển thị; theo dõi Serial Monitor hoặc RFC2217 tại localhost:4003.",
            "Mở dashboard tại http://localhost:8501 để kiểm tra metric và dữ liệu raw.",
            "Tại repo root, chạy docker compose run --rm --build preprocess để cập nhật dữ liệu processed.",
        ],
    )
    add_code(
        document,
        "docker compose up -d --build\n"
        "docker compose ps\n"
        "docker compose logs -f collector\n"
        "docker compose run --rm --build preprocess",
    )
    add_code(
        document,
        "# Chạy ở thư mục firmware\n"
        "pio device monitor -p rfc2217://localhost:4003",
    )

    document.add_heading("10. Kết quả kiểm tra", level=1)
    add_table(
        document,
        ["Hạng mục", "Kết quả ghi nhận"],
        [
            ("PlatformIO", "Build firmware ESP32 thành công."),
            ("MQTT", "Serial có publish=OK; topic telemetry nhận JSON."),
            ("DHT22", "Đã quan sát payload hợp lệ 25.00 °C / 50.00% và sau đó 0.20 °C / 18.50%."),
            ("LDR / HC-SR04", "Payload kiểm tra ghi nhận 499.6 lux và khoảng cách xấp xỉ 100.86 cm."),
            ("InfluxDB", "Truy vấn gần nhất trả temperature=0.2 và humidity=18.5 trong sensor_raw."),
            ("Preprocess", "Một lần chạy thành công đã xử lý 118 cửa sổ."),
            ("Dashboard", "Dịch vụ khởi động và HTTP localhost:8501 trả mã 200."),
        ],
        [3.4, 11.6],
    )
    add_paragraph(
        document,
        "Các số liệu trên là kiểm tra chức năng trong một phiên mô phỏng; chúng không đại diện "
        "cho benchmark tải hoặc độ trễ dài hạn. Dữ liệu Wokwi thay đổi khi người dùng điều chỉnh "
        "thuộc tính cảm biến.",
    )

    document.add_heading("11. Sự cố và cách xử lý", level=1)
    add_table(
        document,
        ["Hiện tượng", "Nguyên nhân / cách kiểm tra"],
        [
            ("Wokwi không tìm thấy wokwi.toml", "Mở đúng firmware làm workspace; file wokwi.toml phải cạnh platformio.ini và đường dẫn firmware phải tương đối với file đó."),
            ("Serial RFC2217 trống", "Đảm bảo simulation đang chạy và tab Wokwi hiển thị; kết nối localhost:4003 ở 115200 baud."),
            ("DHT có giá trị null", "Xem log DHT22 read failed; kiểm tra dây SDA → GPIO15, VCC, GND và payload raw."),
            ("Preprocess KeyError khi thiếu DHT", "Script hiện xử lý field khả dụng, cảnh báo field DHT thiếu và không tự sinh số đo giả."),
            ("MQTT không kết nối", "Kiểm tra broker đang chạy, host IPv4 hiện tại, cổng 1883, firewall và Wokwi Private IoT Gateway."),
            ("Dashboard cũ/không cập nhật", "F5 hoặc Ctrl+F5; xác nhận collector ghi mẫu mới vào sensor_raw."),
        ],
        [4.3, 10.7],
    )

    document.add_heading("12. An toàn và giới hạn", level=1)
    add_paragraph(
        document,
        "Mosquitto hiện cho anonymous access và MQTT cổng 1883 không mã hóa. Cấu hình phù hợp "
        "cho demo trên máy cá nhân, không dùng trực tiếp trong môi trường thật. Triển khai thực "
        "tế cần xác thực, ACL, TLS, secrets ngoài source control và chính sách retention phù hợp.",
    )
    add_paragraph(
        document,
        "Wokwi có thể pause khi tab simulator bị ẩn; điều này ảnh hưởng tính liên tục của chuỗi "
        "thời gian. Giá trị sensor null cần được giữ như missing trong raw để phân biệt với số đo "
        "thật; nội suy chỉ được thực hiện ở bước preprocessing.",
    )

    document.add_heading("13. Kết luận", level=1)
    add_paragraph(
        document,
        "Bài thực hành đã ghép nối được các tầng thiết bị, MQTT, collector Python, InfluxDB, "
        "preprocessing và dashboard. Việc tách raw và processed giúp giữ dữ liệu gốc, kiểm tra "
        "lỗi cảm biến và chạy lại bước xử lý. Hướng phát triển tiếp theo gồm bảo mật MQTT, "
        "retention policy, cập nhật dashboard định kỳ và đánh giá pipeline với nhiều thiết bị.",
    )

    document.add_heading("Tài liệu tham khảo", level=1)
    add_bullets(
        document,
        [
            "Mã nguồn tham khảo Bài thực hành 2 IoT: https://github.com/kareal0907/bai-thuc-hanh-2-IOT",
            "Wokwi documentation: https://docs.wokwi.com/",
            "InfluxDB 2.x documentation: https://docs.influxdata.com/influxdb/v2/",
            "Streamlit documentation: https://docs.streamlit.io/",
            "Tài liệu Bài thực hành số 1 – ESP32, DHT22 và MQTT do sinh viên cung cấp.",
        ],
    )

    document.core_properties.title = "Báo cáo Bài thực hành 2 - IoT"
    document.core_properties.subject = "Thu thập, lưu trữ và tiền xử lý dữ liệu IoT"
    document.core_properties.author = "Vũ Xuân Trường"
    document.save(OUTPUT)
    print(f"Created {OUTPUT}")


if __name__ == "__main__":
    build_report()